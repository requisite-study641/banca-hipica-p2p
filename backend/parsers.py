"""
Parser de cartelera — formatos admitidos en ``POST /admin/races/upload``.

Tipos de archivo admitidos (detección automática por extensión o firma binaria):
    * Texto plano: ``.csv`` / ``.txt`` (UTF-8, cp1252 o latin-1).
    * Excel nativo: ``.xlsx`` / ``.xlsm`` (vía ``openpyxl``) y ``.xls`` legado
      (vía ``xlrd``) — se leen EN MEMORIA y se convierten a un texto CSV
      equivalente (``_read_xlsx`` / ``_read_xls``) para reutilizar, sin
      duplicar código, la misma lógica de extracción de carreras/ejemplares.

Formato A · REPORTE OFICIAL INH (p. ej. "Ejemplares inscritos reunión 39.csv"
o ".xlsx"):
    Diseño de reporte institucional: metadatos por carrera ("Carrera Nro: X",
    "Fecha: dd/mm/aaaa", "Hora: hh:mm" — admite etiqueta y valor en celdas o
    filas separadas) y bloques tabulares de ejemplares con columnas ``N°`` y
    ``Ejemplar`` (alias: ``Nro``, ``Nº``, ``Caballo``, ``Nombre``…).

Formato B · CSV/EXCEL PLANO (una fila por ejemplar):
    Encabezado único en la fila 1 — ``numero_carrera``, ``numero_caballo``,
    ``nombre_caballo``, ``fecha``, ``hora`` (se admiten sinónimos español/
    inglés y variaciones de mayúsculas/minúsculas, ver ``ALIASES``).

Ambos formatos se procesan de forma ATÓMICA: si hay algún error no se importa
ninguna fila y se devuelve el detalle de los errores.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime, time as dt_time
from typing import Dict, List, Optional, Tuple

# ─────────────────────── Librerías de Excel (opcionales) ───────────────────────
# Se importan de forma segura: si no están instaladas en el entorno, el backend
# sigue funcionando normalmente para CSV/texto y solo informa un error claro
# (HTTP 400) si el admin intenta subir un .xlsx/.xls sin la librería disponible.
try:
    import openpyxl  # noqa: F401  (xlsx / xlsm)
except ImportError:  # pragma: no cover
    openpyxl = None

try:
    import xlrd  # noqa: F401  (xls legado)
except ImportError:  # pragma: no cover
    xlrd = None

# ───────────────────────────── Codificaciones ─────────────────────────────
# utf-8-sig preserva la plantilla existente; cp1252/latin-1 cubren el reporte
# INH (latin-1 es el fallback final: decodifica cualquier byte, nunca falla).
ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")

# Extensiones de Excel admitidas (además de .csv / texto plano).
EXCEL_EXTENSIONS = (".xlsx", ".xlsm", ".xls")

FORMAT_LABELS = {
    "INH": "Reporte oficial INH",
    "CSV": "CSV plano (estándar)",
}

# ───────────────────────── Formatos de fecha/hora ─────────────────────────
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y")
TIME_FORMATS = ("%H:%M", "%H:%M:%S", "%H.%M")

# ─────────────────────── Expresiones del reporte INH ───────────────────────
# "Carrera Nro: 1", "CARRERA 1", "Nro. Carrera: 3", "  1) Carrera 2", …
RE_CARRERA = re.compile(
    r"^\s*(?:\d+\s*[.):\-\t]+\s*)?"
    r"(?:carrera|n(?:ro|ú|u)?\.?\s*[°º]?\s*carrera)\b"
    r"\s*(?:n(?:ro|ú|u)?\.?\s*[°º#]?|#)?\s*[:.\-]?\s*(\d{1,3})\b",
    re.IGNORECASE,
)
# "Fecha: 06/10/2026" / "Fecha : 2026-10-06" (estricta: exige el separador).
RE_FECHA = re.compile(
    r"\bfecha\b\s*[:.\-]?\s*([0-9]{1,4}[/-][0-9]{1,2}[/-][0-9]{1,4})", re.IGNORECASE
)
# "Hora: 14:30" / "Hora 14.30".
RE_HORA = re.compile(r"\bhora\b\s*[:.\-]?\s*([0-9]{1,2}[:.][0-9]{2})", re.IGNORECASE)
# Variante textual: "14 h 30", "14 horas 30".
RE_HORA_ALT = re.compile(
    r"\b([0-9]{1,2})\s*(?:h|hs|horas)\s*[:.]?\s*([0-9]{2})\b", re.IGNORECASE
)
# ───────── Etiquetas "sueltas" (sin valor en la misma línea/celda) ─────────
# Algunos exportadores de Excel (celdas combinadas → CSV) dejan la etiqueta
# en una fila y el valor en la fila siguiente, p. ej.:
#   "Carrera Nro:"   (fila N)
#   "1"              (fila N+1)
# Estas expresiones detectan la etiqueta SIN valor, para esperar el dato en
# la siguiente línea no vacía (ver estado ``pending_label`` en ``parse_inh_report``).
RE_LABEL_CARRERA = re.compile(
    r"^\s*(?:carrera|n(?:ro|ú|u)?\.?\s*[°º]?\s*carrera)\b"
    r"\s*(?:n(?:ro|ú|u)?\.?\s*[°º#]?|#)?\s*[:.\-]?\s*$",
    re.IGNORECASE,
)
RE_LABEL_FECHA = re.compile(r"^\s*fecha\s*[:.\-]?\s*$", re.IGNORECASE)
RE_LABEL_HORA = re.compile(r"^\s*hora\s*[:.\-]?\s*$", re.IGNORECASE)
# Valores "sueltos" esperados tras una etiqueta pendiente.
RE_VALUE_NUM = re.compile(r"^\s*(\d{1,3})\s*$")
RE_VALUE_FECHA = re.compile(r"^\s*([0-9]{1,4}[/-][0-9]{1,2}[/-][0-9]{1,4})\s*$")
RE_VALUE_HORA = re.compile(r"^\s*([0-9]{1,2}[:.][0-9]{2})\s*$")
# Detección de formato: cabeceras de carrera o tablas institucionales.
RE_INH_CARRERA = re.compile(
    r"^\s*(?:\d+\s*[.):\-\t]+\s*)?carrera\b", re.IGNORECASE | re.MULTILINE
)
RE_INH_TABLA = re.compile(r"n[°º]\s*[;|\t]", re.IGNORECASE)

# ──────────────────── Columnas de la tabla de ejemplares ────────────────────
# "dorsal" es el alias usado por algunos bloques del reporte oficial real
# en vez de "N°" para la columna del número de ejemplar.
NUMBER_COLS = {"n", "no", "nro", "nros", "num", "numero", "dorsal"}
NAME_COLS = {"ejemplar", "ejemplares", "caballo", "nombre", "horse", "horse_name"}
DELIMITERS = (";", ",", "\t", "|")


# ────────────────────────────── Utilidades ──────────────────────────────
def decode_upload(raw: bytes) -> Tuple[str, str]:
    """Decodifica el archivo: utf-8-sig → cp1252 → latin-1. Devuelve (texto, codificación)."""
    for encoding in ENCODINGS:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace"), "latin-1"  # pragma: no cover


def _norm_cell(cell: str) -> str:
    """'N°' → 'n', 'Ejemplar;' → 'ejemplar' (sin tildes ni símbolos)."""
    text = unicodedata.normalize("NFKD", str(cell or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def _norm_header(header: Optional[str]) -> str:
    return (header or "").strip().lower().replace(" ", "_")


def _split_row(line: str, delimiter: str) -> List[str]:
    try:
        return next(csv.reader([line], delimiter=delimiter))
    except (csv.Error, StopIteration):  # pragma: no cover
        return [line]


def _parse_int(value: str) -> int:
    """Convierte '1', '01', '1.', ' 2 '… en entero; lanza ValueError si no hay dígitos."""
    match = re.search(r"\d+", str(value))
    if not match:
        raise ValueError(f"número no reconocido: '{value}'")
    return int(match.group())


def _parse_datetime(fecha: str, hora: Optional[str]) -> datetime:
    date_part = None
    for fmt in DATE_FORMATS:
        try:
            date_part = datetime.strptime(fecha.strip(), fmt)
            break
        except ValueError:
            continue
    if date_part is None:
        raise ValueError(f"fecha no reconocida: '{fecha}' (use AAAA-MM-DD o DD/MM/AAAA)")

    parsed_time = None
    for fmt in TIME_FORMATS:
        try:
            parsed_time = datetime.strptime((hora or "00:00").strip(), fmt).time()
            break
        except ValueError:
            continue
    if parsed_time is None:
        raise ValueError(f"hora no reconocida: '{hora}' (use HH:MM)")

    # REGLA AUTOMÁTICA HÍPICA: Las carreras son vespertinas (PM).
    # Si la hora viene en formato de 1:00 a 11:59 (sin especificar formato 24h),
    # se interpreta automáticamente como horario de la tarde (sumando 12 horas).
    hour = parsed_time.hour
    minute = parsed_time.minute
    if 1 <= hour <= 11:
        hour += 12

    final_time = dt_time(hour, minute)
    return datetime.combine(date_part.date(), final_time)


# ───────────────────── Detección de cabecera de tabla ─────────────────────
def _detect_header_row(cells: List[str]) -> Optional[Tuple[str, int, int]]:
    """
    Devuelve (delimitador, índice N°, índice Ejemplar) de una fila csv.

    La validación se basa en las columnas (N°, Ejemplar, etc.) ya extraídas
    por ``csv.reader`` con el delimitador correcto. El delimitador devuelto
    es siempre el usado por el parser principal (``;``).
    """
    idx_num = idx_name = None
    for i, cell in enumerate(cells):
        key = _norm_cell(cell)
        if idx_num is None and key in NUMBER_COLS:
            idx_num = i
        if idx_name is None and (key in NAME_COLS or key.startswith("ejemplar")):
            idx_name = i
    if idx_num is not None and idx_name is not None and idx_num != idx_name:
        return DELIMITERS[0], idx_num, idx_name
    return None


def _detect_header(line: str) -> Optional[Tuple[str, int, int]]:
    """Devuelve (delimitador, índice N°, índice Ejemplar) si la línea es cabecera."""
    for delimiter in DELIMITERS:
        if delimiter not in line:
            continue
        cells = _split_row(line, delimiter)
        idx_num = idx_name = None
        for i, cell in enumerate(cells):
            key = _norm_cell(cell)
            if idx_num is None and key in NUMBER_COLS:
                idx_num = i
            if idx_name is None and (key in NAME_COLS or key.startswith("ejemplar")):
                idx_name = i
        if idx_num is not None and idx_name is not None and idx_num != idx_name:
            return delimiter, idx_num, idx_name
    return None


# ───────────────────────── Reporte oficial INH ─────────────────────────
def _clean_cell_text(value: object) -> str:
    """Une saltos de línea internos de una celda (reclamo/precio en varias
    líneas, celdas combinadas de Excel) en un solo espacio y recorta bordes."""
    text = str(value) if value is not None else ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return " ".join(part.strip() for part in text.split("\n") if part.strip()).strip()


def parse_inh_report(text):
    """
    Parsea el reporte oficial INH leyendo TODO el archivo con ``csv.reader``
    (delimitador ``;``, ``quotechar='"'``) para que los campos multilínea
    citados (``"Carrera Nro:\\n1"``, nombres de ejemplares de reclamo con
    precio en varias líneas, etc.) no fragmenten los bloques.

    Soporta además el caso en que la etiqueta (``Carrera Nro:`` / ``Fecha:`` /
    ``Hora:``) queda en una fila y su valor en la fila siguiente (exportación
    de celdas combinadas de Excel a CSV), mediante el estado ``pending_label``.
    """
    rows_by_key = {}
    errors = []
    races = {}
    report_fecha = None
    report_hora = None

    current_race = None
    in_table = False
    skip_block = False
    pending_suplente = False
    # Etiquetas "sueltas" por COLUMNA: el reporte real del INH organiza los
    # metadatos en una cuadrícula donde cada columna puede tener su propia
    # etiqueta ("Carrera Nro:", "Fecha:", "Hora:"...) con el valor en la
    # MISMA columna de la fila siguiente (celdas combinadas de Excel). Se
    # procesa CELDA POR CELDA (no la fila entera concatenada) porque una
    # misma fila puede traer varias etiquetas/valores en columnas distintas.
    pending_labels: Dict[int, str] = {}
    idx_num = None
    idx_name = None

    reader = csv.reader(io.StringIO(text), delimiter=";", quotechar='"')
    for raw_cells in reader:
        line_no = reader.line_num
        # Limpia saltos de línea internos de cada celda (campos multilínea).
        cells = [_clean_cell_text(c) for c in raw_cells]

        if not cells or all(c == "" for c in cells):
            in_table = False
            pending_suplente = False
            continue

        if any(_norm_cell(cell).startswith("suplente") for cell in cells):
            pending_suplente = True
            pending_labels = {}
            continue

        header = _detect_header_row(cells)
        if header:
            _delimiter, idx_num, idx_name = header
            skip_block = pending_suplente
            in_table = not skip_block
            pending_labels = {}
            if in_table and current_race is None:
                errors.append(f"Línea {line_no}: tabla de ejemplares sin bloque.")
            continue

        # ── Filas de la tabla de ejemplares: SOLO si la columna N°/dorsal es
        # puramente numérica (evita que pies de página/observaciones largas
        # que contienen algún dígito suelto se confundan con un ejemplar).
        if in_table and current_race is not None and not skip_block and idx_num is not None and idx_name is not None:
            if len(cells) > max(idx_num, idx_name):
                number_raw = cells[idx_num].strip()
                horse_name = cells[idx_name].strip()
                if number_raw or horse_name:
                    if re.fullmatch(r"\d{1,3}\.?", number_raw):
                        horse_number = int(number_raw.rstrip("."))
                        if not horse_name:
                            errors.append(
                                f"Línea {line_no}: el ejemplar N.{horse_number} no tiene nombre."
                            )
                        else:
                            rows_by_key[(current_race, horse_number)] = {
                                "race_number": current_race,
                                "horse_number": horse_number,
                                "horse_name": horse_name,
                            }
                        continue
                    # Hay texto en la columna N° pero no es un número puro
                    # (pie de página/observación institucional) o no hay
                    # número en esa columna → se ignora sin generar error
                    # falso y sin cortar la tabla.
                    continue

        # ── Metadatos: se evalúa CADA celda de la fila de forma independiente.
        for col_idx, raw_cell in enumerate(cells):
            cell = raw_cell.strip()
            if not cell:
                continue

            # ── Valor "suelto" esperado tras una etiqueta pendiente EN ESTA
            # MISMA COLUMNA (celda previa sin el dato inline: "Carrera Nro:"
            # / "Fecha:" / "Hora:" en una fila, valor en la fila siguiente).
            pending = pending_labels.get(col_idx)
            if pending == "carrera":
                value_match = RE_VALUE_NUM.match(cell)
                if value_match:
                    current_race = int(value_match.group(1))
                    in_table = skip_block = False
                    pending_suplente = False
                    races.setdefault(current_race, {"fecha": None, "hora": None})
                    del pending_labels[col_idx]
                    continue
                del pending_labels[col_idx]
            elif pending == "fecha":
                value_match = RE_VALUE_FECHA.match(cell)
                if value_match:
                    value = value_match.group(1)
                    if current_race is not None:
                        races.setdefault(current_race, {"fecha": None, "hora": None})
                        if races[current_race]["fecha"] is None:
                            races[current_race]["fecha"] = value
                    elif report_fecha is None:
                        report_fecha = value
                    del pending_labels[col_idx]
                    continue
                del pending_labels[col_idx]
            elif pending == "hora":
                value_match = RE_VALUE_HORA.match(cell)
                if value_match:
                    value = value_match.group(1)
                    if current_race is not None:
                        races.setdefault(current_race, {"fecha": None, "hora": None})
                        if races[current_race]["hora"] is None:
                            races[current_race]["hora"] = value
                    elif report_hora is None:
                        report_hora = value
                    del pending_labels[col_idx]
                    continue
                del pending_labels[col_idx]

            match = RE_CARRERA.match(cell)
            if match:
                current_race = int(match.group(1))
                in_table = skip_block = False
                pending_suplente = False
                races.setdefault(current_race, {"fecha": None, "hora": None})
                continue
            if RE_LABEL_CARRERA.match(cell):
                # Etiqueta sin número: el valor llega en la misma columna de
                # la(s) siguiente(s) fila(s).
                pending_labels[col_idx] = "carrera"
                continue

            match_fecha = RE_FECHA.search(cell)
            if match_fecha:
                value = match_fecha.group(1)
                if current_race is not None:
                    races.setdefault(current_race, {"fecha": None, "hora": None})
                    if races[current_race]["fecha"] is None:
                        races[current_race]["fecha"] = value
                elif report_fecha is None:
                    report_fecha = value
                continue
            if RE_LABEL_FECHA.match(cell):
                pending_labels[col_idx] = "fecha"
                continue

            match_hora = RE_HORA.search(cell)
            hora_value = match_hora.group(1) if match_hora else None
            if hora_value is None:
                match_alt = RE_HORA_ALT.search(cell)
                if match_alt:
                    hora_value = f"{match_alt.group(1)}:{match_alt.group(2)}"
            if hora_value:
                if current_race is not None:
                    races.setdefault(current_race, {"fecha": None, "hora": None})
                    if races[current_race]["hora"] is None:
                        races[current_race]["hora"] = hora_value
                elif report_hora is None:
                    report_hora = hora_value
                continue
            if RE_LABEL_HORA.match(cell):
                pending_labels[col_idx] = "hora"
                continue

    datetimes = {}
    for race_number, meta in races.items():
        fecha = meta.get("fecha") or report_fecha
        hora = meta.get("hora") or report_hora
        if not fecha:
            errors.append(f"Carrera {race_number}: falta la fecha.")
            continue
        if not hora:
            errors.append(f"Carrera {race_number}: falta la hora.")
            continue
        try:
            datetimes[race_number] = _parse_datetime(fecha, hora)
        except ValueError as exc:
            errors.append(f"Carrera {race_number}: {exc}")

    races_with_rows = {key[0] for key in rows_by_key}
    for race_number in sorted(races):
        if race_number not in races_with_rows:
            errors.append(f"Carrera {race_number}: no tiene ejemplares en su tabla.")

    rows = []
    for (race_number, _horse), data in sorted(rows_by_key.items()):
        race_datetime = datetimes.get(race_number)
        if race_datetime is None:
            continue
        rows.append({**data, "race_datetime": race_datetime})

    return {"format": "INH", "rows": rows, "errors": errors}



# ─────────────────────────── CSV plano (UTF-8) ───────────────────────────
# Encabezados aceptados por campo (minúsculas, sin espacios).
ALIASES: Dict[str, Tuple[str, ...]] = {
    "horse_number": (
        "numero_caballo", "num_caballo", "caballo", "caballo_nro", "caballo_numero",
        "horse_number", "no_caballo", "nro_caballo", "n_caballo",
    ),
    "horse_name": ("nombre_caballo", "nombre", "horse_name", "caballo_nombre", "ejemplar"),
    "race_number": (
        "numero_carrera", "num_carrera", "carrera", "carrera_nro", "carrera_numero",
        "race_number", "nro_carrera", "n_carrera",
    ),
    "fecha": ("fecha", "fecha_carrera", "date", "race_date"),
    "hora": ("hora", "hora_carrera", "time", "race_time"),
}
EXPECTED_HEADER = "numero_caballo,nombre_caballo,numero_carrera,fecha,hora"
# Conjunto de TODOS los alias reconocidos (para la detección temprana de
# Formato B en ``detect_format``, evitando que una cabecera de columna que
# empiece con la palabra "carrera" se confunda con un bloque INH real).
ALL_HEADER_ALIASES = {alias for aliases in ALIASES.values() for alias in aliases}


def parse_flat_csv(text: str) -> dict:
    """
    Parser del CSV plano clásico (una fila por ejemplar).

    Se utiliza ``csv.DictReader`` para manejar correctamente:
      * comillas dentro de campos (``"texto ""cita"" texto"``),
      * campos que contienen el delimitador o saltos de línea,
      * codificaciones UTF-8 / UTF-8 con BOM.
    """
    rows: List[dict] = []
    errors: List[str] = []

    first_line = next((ln for ln in text.splitlines() if ln.strip()), "")
    delimiter = max(DELIMITERS, key=first_line.count) if first_line else ","
    if first_line.count(delimiter) == 0:
        delimiter = ","

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        return {"format": "CSV", "rows": [], "errors": ["El CSV no tiene encabezados."]}

    headers = {_norm_header(h) for h in reader.fieldnames}
    missing = [
        label
        for field, label in (
            ("horse_number", "numero_caballo"),
            ("horse_name", "nombre_caballo"),
            ("race_number", "numero_carrera"),
            ("fecha", "fecha"),
            ("hora", "hora"),
        )
        if not any(alias in headers for alias in ALIASES[field])
    ]
    if missing:
        return {
            "format": "CSV",
            "rows": [],
            "errors": [
                f"Faltan columnas: {', '.join(missing)}. Encabezado esperado: {EXPECTED_HEADER}"
            ],
        }

    def pick(row: Dict[str, str], field: str) -> Optional[str]:
        for alias in ALIASES[field]:
            if alias in row:
                cleaned = _clean_cell_text(row[alias])
                if cleaned:
                    return cleaned
        return None

    for index, raw_row in enumerate(reader, start=2):  # La fila 1 es el encabezado.
        row = {_norm_header(k): v for k, v in raw_row.items() if k is not None}
        horse_number = pick(row, "horse_number")
        horse_name = pick(row, "horse_name")
        race_number = pick(row, "race_number")
        fecha = pick(row, "fecha")
        hora = pick(row, "hora")

        if not any([horse_number, horse_name, race_number]):
            # Fila en blanco (o artefacto de formato heredado de Excel: una
            # celda de fecha "fantasma" en una fila vacía sin ningún otro dato).
            continue
        try:
            if not horse_number:
                raise ValueError("falta numero_caballo")
            if not horse_name:
                raise ValueError("falta nombre_caballo")
            if not race_number:
                raise ValueError("falta numero_carrera")
            if not fecha:
                raise ValueError("falta fecha")
            rows.append(
                {
                    "race_number": _parse_int(race_number),
                    "horse_number": _parse_int(horse_number),
                    "horse_name": horse_name,
                    "race_datetime": _parse_datetime(fecha, hora),
                }
            )
        except (ValueError, TypeError) as exc:
            errors.append(f"Fila {index}: {exc}")

    if not rows and not errors:
        errors.append("El archivo no contiene filas de datos.")
    return {"format": "CSV", "rows": rows, "errors": errors}


# ───────────────────────── Lectura nativa de Excel ─────────────────────────
# El archivo Excel se convierte en memoria a un texto CSV equivalente
# (delimitador ``;``, citado con comillas dobles) para REUTILIZAR sin
# duplicar la lógica de extracción de carreras/ejemplares existente
# (``parse_inh_report`` / ``parse_flat_csv`` / ``detect_format``).
EXCEL_SIGNATURES = (
    (b"PK\x03\x04", ".xlsx"),        # ZIP (xlsx/xlsm modernos)
    (b"\xd0\xcf\x11\xe0", ".xls"),   # OLE2 (xls legado)
)


def _sniff_excel_extension(raw: bytes) -> Optional[str]:
    """Detecta por firma binaria si el archivo es Excel, aunque el nombre
    subido no tenga extensión o el navegador no la haya conservado."""
    for signature, extension in EXCEL_SIGNATURES:
        if raw.startswith(signature):
            return extension
    return None


def _cell_to_str(value: object) -> str:
    """Convierte el valor nativo de una celda de Excel a texto, preservando
    el formato de fecha/hora esperado por el parser (dd/mm/aaaa, HH:MM)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() in (dt_time(0, 0, 0),):
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, dt_time):
        return value.strftime("%H:%M")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def _trim_blank_edges(rows: List[List[str]]) -> List[List[str]]:
    """Elimina filas COMPLETAMENTE vacías al inicio y al final de la matriz
    (artefacto típico de Excel: fila 1 en blanco antes del encabezado real, o
    filas finales con solo formato heredado). Las filas vacías intermedias
    se conservan porque delimitan bloques en el reporte INH."""

    def is_blank(row: List[str]) -> bool:
        return not any(str(cell or "").strip() for cell in row)

    start = 0
    while start < len(rows) and is_blank(rows[start]):
        start += 1
    end = len(rows)
    while end > start and is_blank(rows[end - 1]):
        end -= 1
    return rows[start:end]


def _rows_to_csv_text(rows: List[List[str]]) -> str:
    """Serializa una matriz de celdas ya convertidas a texto, como CSV
    delimitado por ``;`` y citado con comillas dobles (compatible con
    ``csv.reader``/``csv.DictReader`` usados por el resto del parser)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quotechar='"', quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
    for row in _trim_blank_edges(rows):
        writer.writerow(row)
    return buffer.getvalue()


def _read_xlsx(raw: bytes) -> str:
    """Lee la primera hoja de un .xlsx/.xlsm con ``openpyxl`` y la convierte
    a texto CSV equivalente."""
    if openpyxl is None:
        raise ValueError(
            "El servidor no tiene instalada la librería 'openpyxl' para leer "
            "archivos .xlsx. Instálala con 'pip install openpyxl' o sube el "
            "archivo en formato CSV."
        )
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as exc:  # noqa: BLE001 — cualquier error de lectura se traduce a 400.
        raise ValueError(f"No se pudo leer el archivo .xlsx: {exc}")
    sheet = workbook.active
    rows = [[_cell_to_str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    return _rows_to_csv_text(rows)


def _read_xls(raw: bytes) -> str:
    """Lee la primera hoja de un .xls (formato legado) con ``xlrd`` y la
    convierte a texto CSV equivalente."""
    if xlrd is None:
        raise ValueError(
            "El servidor no tiene instalada la librería 'xlrd' para leer "
            "archivos .xls (formato antiguo). Instálala con 'pip install "
            "xlrd' o sube el archivo en formato .xlsx o CSV."
        )
    try:
        book = xlrd.open_workbook(file_contents=raw)
        sheet = book.sheet_by_index(0)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"No se pudo leer el archivo .xls: {exc}")

    rows: List[List[str]] = []
    for row_idx in range(sheet.nrows):
        row_values = []
        for col_idx in range(sheet.ncols):
            cell = sheet.cell(row_idx, col_idx)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                except Exception:  # noqa: BLE001 — si falla, se deja el valor crudo.
                    pass
            row_values.append(_cell_to_str(value))
        rows.append(row_values)
    return _rows_to_csv_text(rows)


# ───────────────────────── Detección de formato ─────────────────────────
def detect_format(text: str) -> str:
    """'INH' si el archivo tiene marcadores del reporte institucional; si no, 'CSV'."""
    # Prioridad: si la PRIMERA línea no vacía es un encabezado de Formato B
    # reconocido (p. ej. "carrera,caballo_nro,nombre,fecha,hora"), se trata
    # como CSV plano de inmediato — evita que la palabra "carrera" al inicio
    # de la cabecera de columna se confunda con un bloque "Carrera Nro: X"
    # del reporte INH (``RE_INH_CARRERA`` solo exige que la línea EMPIECE con
    # "carrera", sin exigir separador/número, por diseño permisivo).
    first_line = next((ln for ln in text.splitlines() if ln.strip()), "")
    if first_line:
        delimiter = max(DELIMITERS, key=first_line.count)
        if first_line.count(delimiter) > 0:
            tokens = {_norm_header(tok) for tok in first_line.split(delimiter)}
            has_race = any(alias in tokens for alias in ALIASES["race_number"])
            has_horse_num = any(alias in tokens for alias in ALIASES["horse_number"])
            has_horse_name = any(alias in tokens for alias in ALIASES["horse_name"])
            if has_race and has_horse_num and has_horse_name:
                return "CSV"

    if RE_INH_CARRERA.search(text):
        return "INH"
    lowered = text.lower()
    if "ejemplare" in lowered and (RE_INH_TABLA.search(text) or "inscritos" in lowered):
        return "INH"
    return "CSV"


# ─────────────────────────── Fachada principal ───────────────────────────
def parse_cartelera(raw: bytes, filename: Optional[str] = None) -> dict:
    """
    Decodifica (texto CSV o Excel nativo), detecta el formato (INH vs CSV
    plano) y parsea el archivo.

    * Si ``filename`` termina en ``.xlsx``/``.xlsm``/``.xls`` (o el contenido
      coincide con la firma binaria de esos formatos aunque no se indique
      extensión), el archivo se lee EN MEMORIA con ``openpyxl``/``xlrd`` y se
      convierte a un texto CSV equivalente (delimitador ``;``) que reutiliza
      exactamente la misma lógica de extracción de carreras/ejemplares.
    * En cualquier otro caso se trata como texto (CSV plano o reporte INH).

    Devuelve ``{"format", "format_label", "encoding", "rows", "errors"}`` donde
    cada fila es ``{race_number, horse_number, horse_name, race_datetime}``.
    Lanza ``ValueError`` si el archivo está vacío o no se puede leer.
    """
    if not raw:
        raise ValueError("El archivo está vacío.")

    extension = (filename or "").strip().lower()
    extension = extension[extension.rfind("."):] if "." in extension else ""
    sniffed = _sniff_excel_extension(raw)

    if extension in (".xlsx", ".xlsm") or (not extension and sniffed == ".xlsx"):
        text = _read_xlsx(raw)
        encoding = "excel (.xlsx)"
    elif extension == ".xls" or (not extension and sniffed == ".xls"):
        text = _read_xls(raw)
        encoding = "excel (.xls)"
    else:
        text, encoding = decode_upload(raw)
        text = text.lstrip("\ufeff").lstrip("ï»¿")

    if not text.strip():
        raise ValueError("El archivo está vacío.")

    result = parse_inh_report(text) if detect_format(text) == "INH" else parse_flat_csv(text)
    result["encoding"] = encoding
    result["format_label"] = FORMAT_LABELS[result["format"]]
    return result



